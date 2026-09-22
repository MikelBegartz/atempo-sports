"""Importación sencilla de listas: equipos, personas y plantillas (CSV/Excel)."""

from __future__ import annotations

import csv
import io
import unicodedata
from dataclasses import dataclass, field

from sqlalchemy.orm import Session

from app.db import Person, Team, TeamMembership
from app.teams_meta import infer_branch

XLSX_MAGIC = b"PK\x03\x04"
XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

# Capçaleres conegudes (normalitzades: minúscules, sense espais extrems)
PERSON_HEADERS = {
    "persona", "person", "nombre", "name", "nome", "nom",
    "nom i cognom", "nom i cognoms", "nom complet",
    "jugadors", "jugadores",
}
TEAM_HEADERS = {"equipo", "team", "equip", "equips", "equipa", "squadra", "équipe", "teams"}
COACH_HEADERS = {
    "entrenador", "entrenadora", "entrenadors", "entrenadores",
    "treinador", "treinadora", "coach",
}
ROLE_HEADERS = {"rol", "role", "função", "ruolo"}
OTHER_HEADERS = {"jugador", "player", "categoria", "category", "dorsal"}
ALL_HEADERS = PERSON_HEADERS | TEAM_HEADERS | COACH_HEADERS | ROLE_HEADERS | OTHER_HEADERS

# Paraules típiques en noms d'equip (per a fulls sense capçalera)
TEAM_WORDS = {
    "benjamí", "benjami", "prebenjamí", "prebenjami", "aleví", "alevi",
    "infantil", "juvenil", "senior", "sènior", "cadet", "iniciació",
    "iniciacio", "femení", "femeni", "masculí", "masculi", "promeses",
}


ROLE_MAP = {
    "player": "player",
    "jugador": "player",
    "jogadora": "player",
    "jogadora/jogador": "player",
    "jogador": "player",
    "giocatore": "player",
    "joueur": "player",
    "spieler": "player",
    "coach": "coach",
    "entrenador": "coach",
    "entrenadora": "coach",
    "treinador": "coach",
    "allenatore": "coach",
    "entraîneur": "coach",
    "trainer": "coach",
    "reinforce": "reinforce",
    "refuerzo": "reinforce",
    "reforç": "reinforce",
    "reforço": "reinforce",
    "rinforzo": "reinforce",
    "renfort": "reinforce",
    "verstärkung": "reinforce",
}


@dataclass
class ImportReport:
    teams_created: int = 0
    people_created: int = 0
    links_created: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors or (
            self.teams_created + self.people_created + self.links_created > 0
        )


def _detect_delimiter(sample: str) -> str:
    """Detecta el separador mirando las primeras líneas (; , o tabulador)."""
    lines = [l for l in sample.splitlines() if l.strip()][:5]
    counts = {";": 0, ",": 0, "\t": 0}
    for line in lines:
        for d in counts:
            counts[d] += line.count(d)
    best = max(counts, key=lambda d: counts[d])
    return best if counts[best] > 0 else ";"


def _norm(value: str | None) -> str:
    return (value or "").strip()


_ROLE_WORDS = {
    "jugador", "jugadora", "jugadores", "jugadoras", "jogador", "jogadora",
    "jugadors", "jugadores", "jugador/a", "jugsdor", "jugsdora",
    "jugadorr", "jugdor", "jugaador",
    "entrenador", "entrenadora", "entrenadors", "entrenadores",
    "entrenador/a", "treinador", "treinadora",
    "delegado", "delegada", "delegat", "delegats", "delegades",
    "portero", "portera", "porter", "porteros", "porteras",
    "player", "coach", "jugadora/o",
}


def canon_person_name(name: str | None) -> str:
    """Forma canònica d'un nom: sense comes, sense paraules de rol
    (jugador/a, entrenador/a, delegat...), espais col·lapsats.
    "ÈRIK, REY, JUGADOR" == "ÈRIK,REY" == "ÈRIK REY"."""
    s = " ".join((name or "").replace(",", " ").split())
    words = [w for w in s.split() if w.casefold() not in _ROLE_WORDS]
    return " ".join(words)


def person_role_flags(raw: str) -> tuple[bool, bool]:
    """Detecta si el text portava rol enganxat: (entrenador, delegat)."""
    low = (raw or "").casefold()
    return ("entrenador" in low or "entrenadora" in low), ("delegad" in low)


def _strip_accents(s: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c)
    )


def person_key(name: str | None) -> str:
    """Clau de deduplicació: nom canònic, sense accents ni majúscules.
    "ÀVILA, LLUC" i "avila lluc" són la mateixa persona."""
    return _strip_accents(canon_person_name(name)).casefold()


def find_person_canon(db: Session, season_id: int, name: str) -> Person | None:
    """Cerca persona per clau canònica (espais, comes, rol, accents,
    majúscules — tot ignorat)."""
    key = person_key(name)
    if not key:
        return None
    for p in db.query(Person).filter(Person.season_id == season_id).all():
        if person_key(p.full_name) == key:
            return p
    return None


def _norm_role(value: str) -> str:
    key = _norm(value).casefold()
    return ROLE_MAP.get(key, "player")


def _truthy(value: str) -> bool:
    return _norm(value).casefold() in {"1", "true", "yes", "si", "sí", "x", "s"}


def _rows_from_table(header: list[str], data_rows: list[list[str]]) -> list[dict[str, str]]:
    """Normaliza cabeceras (minúsculas, sin espacios) y devuelve dicts por fila."""
    field_map = [_norm(h).casefold() for h in header]
    rows: list[dict[str, str]] = []
    for raw_row in data_rows:
        clean = {}
        for i, key in enumerate(field_map):
            if not key:
                continue
            clean[key] = _norm(raw_row[i] if i < len(raw_row) else "")
        if any(clean.values()):
            rows.append(clean)
    return rows


def parse_csv_table(raw: str) -> list[list[str]]:
    """Texto pegado/CSV → tabla de celdas (separador ; , o tabulador)."""
    text = (raw or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return []
    # Quitar BOM Excel
    if text.startswith("\ufeff"):
        text = text[1:]
    delim = _detect_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    return [list(r) for r in reader]


def parse_csv_text(raw: str, kind: str = "roster") -> list[dict[str, str]]:
    """Si hay fila de cabeceras reconocida la usa; si no, modo posicional
    según el formulario (teams/people/roster)."""
    table = parse_csv_table(raw)
    if not table:
        return []
    hi = _find_header_row(table)
    if hi is None:
        return headerless_kind_rows(kind, table)
    return _rows_from_table(table[hi], table[hi + 1:])


def decode_upload_text(content: bytes) -> str:
    """Decodifica un CSV/texto probando UTF-8 (con BOM) y luego cp1252 (Excel ES)."""
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _sheet_table(ws) -> list[list[str]]:
    """Devuelve la hoja como tabla de texto, sin filas vacías iniciales/finales."""
    table: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        table.append([_cell_text(c) for c in row])
    while table and not any(v.strip() for v in table[0]):
        table.pop(0)
    while table and not any(v.strip() for v in table[-1]):
        table.pop()
    return table


def _find_header_row(table: list[list[str]]) -> int | None:
    """Busca en las primeras filas la que más celdas coincide con cabeceras
    conocidas. Devuelve su índice o None si ninguna sirve."""
    best_i, best_score = None, 0
    for i, row in enumerate(table[:15]):
        score = sum(1 for c in row if _norm(c).casefold() in ALL_HEADERS)
        if score > best_score:
            best_i, best_score = i, score
    return best_i


def _headerless_rows(sheet_title: str, table: list[list[str]]) -> list[dict[str, str]]:
    """Hoja sin cabeceras (listado tipo federación): detecta la columna de
    nombres (la que más comas tiene, estilo 'COGNOMS, NOM') y la de equipo
    (la que más palabras de categoría contiene)."""
    if not table:
        return []
    ncols = max(len(r) for r in table)

    def col(c: int) -> list[str]:
        return [_norm(r[c]) if c < len(r) else "" for r in table]

    def _teamish(vals: list[str]) -> int:
        return sum(
            1 for v in vals if any(w in v.casefold() for w in TEAM_WORDS)
        )

    name_i, name_score = None, 0
    for c in range(ncols):
        vals = col(c)
        score = sum(1 for v in vals if "," in v)
        if score == 0:
            score = sum(
                1 for v in vals
                if len(v.split()) >= 2
                and v.replace(" ", "").replace("'", "").replace("-", "").isalpha()
            )
        score -= 3 * _teamish(vals)
        if score > name_score:
            name_i, name_score = c, score
    if name_i is None:
        # Cap columna amb cara de nom: queda't la que menys sembli un equip
        name_i = min(range(ncols), key=lambda c: _teamish(col(c)))
    if name_i is None or name_score <= 0 and not any(col(name_i)):
        return []

    team_i, team_key = None, (0, 0)
    for c in range(ncols):
        if c == name_i:
            continue
        vals = [v for v in col(c) if v]
        score = _teamish(vals)
        distinct = len({v.casefold() for v in vals})
        if score > 0 and (score, distinct) > team_key:
            team_i, team_key = c, (score, distinct)

    rows: list[dict[str, str]] = []
    for r in table:
        name = _norm(r[name_i]) if name_i < len(r) else ""
        team = _norm(r[team_i]) if (team_i is not None and team_i < len(r)) else ""
        rows.append({"equipo": team or sheet_title, "persona": name})
    return rows


def headerless_kind_rows(
    kind: str, table: list[list[str]], sheet_title: str = ""
) -> list[dict[str, str]]:
    """Sense capçalera reconeixible: mapeig posicional segons el formulari.
    - teams: cada fila = un equip
    - people/roster: detecció de columna de noms i d'equip"""
    rows = _headerless_rows(sheet_title, table)
    if kind == "teams":
        out = [
            {"equipo": r.get("equipo") or r.get("persona") or ""}
            for r in rows
        ]
        if not any(r["equipo"] for r in out):
            out = [
                {"equipo": next((c for c in r if c.strip()), "")}
                for r in table
            ]
        return out
    if not rows and kind == "people":
        rows = [
            {"persona": next((c for c in r if c.strip()), "")}
            for r in table
        ]
    return rows


def rows_from_excel(content: bytes, kind: str = "roster") -> list[dict[str, str]]:
    """Lee todas las hojas de un .xlsx/.xlsm. En cada hoja busca la fila de
    cabeceras (equipo/persona/rol, NOM I COGNOM/EQUIP/ENTRENADOR, etc.) y, si
    no hay, usa detección posicional para listados sin cabecera."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ValueError("excel_unsupported") from e
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError("excel_unreadable") from e
    rows: list[dict[str, str]] = []
    for ws in wb.worksheets:
        table = _sheet_table(ws)
        if not table:
            continue
        hi = _find_header_row(table)
        if hi is None:
            rows.extend(headerless_kind_rows(kind, table, ws.title or ""))
        else:
            rows.extend(_rows_from_table(table[hi], table[hi + 1:]))
    wb.close()
    return rows


def rows_from_upload(filename: str, content: bytes, kind: str = "roster") -> list[dict[str, str]]:
    """Elige parser según el contenido del fichero (Excel binario o texto/CSV)."""
    if content[:4] == XLSX_MAGIC:
        return rows_from_excel(content, kind)
    if content[:8] == XLS_MAGIC:
        raise ValueError("excel_xls_old")
    name = (filename or "").casefold()
    if name.endswith((".xlsx", ".xlsm")):
        return rows_from_excel(content, kind)
    if name.endswith(".xls"):
        raise ValueError("excel_xls_old")
    return parse_csv_text(decode_upload_text(content), kind)


def _team_key(name: str | None) -> str:
    """Clau de deduplicació d'equips: sense accents ni majúscules.
    'PREBENJAMI B' i 'PREBENJAMÍ B' són el mateix equip."""
    return _strip_accents(" ".join((name or "").split())).casefold()


def _get_or_create_team(db: Session, season_id: int, name: str, category: str | None, report: ImportReport) -> Team | None:
    name = _norm(name)
    if not name:
        return None
    team = (
        db.query(Team)
        .filter(Team.season_id == season_id, Team.name == name)
        .first()
    )
    if not team:
        key = _team_key(name)
        for t in db.query(Team).filter(Team.season_id == season_id).all():
            if _team_key(t.name) == key:
                team = t
                break
    if team:
        return team
    cat = _norm(category) or None
    team = Team(
        season_id=season_id,
        name=name,
        category=cat,
        branch=infer_branch(name, cat),
    )
    db.add(team)
    db.flush()
    report.teams_created += 1
    return team


def _get_or_create_person(
    db: Session,
    season_id: int,
    name: str,
    role: str,
    report: ImportReport,
) -> Person | None:
    name = canon_person_name(name)
    if not name:
        return None
    person = find_person_canon(db, season_id, name)
    if person:
        # Ampliar flags si el CSV trae rol coach/player
        if role == "coach" and not person.is_coach:
            person.is_coach = True
        if role == "player" and not person.is_player:
            person.is_player = True
        return person
    person = Person(
        season_id=season_id,
        full_name=name,
        is_player=role != "coach",
        is_coach=role == "coach",
    )
    db.add(person)
    db.flush()
    report.people_created += 1
    return person


def _first(row: dict[str, str], keys) -> str:
    for k in keys:
        if row.get(k):
            return row[k]
    return ""


def import_roster_rows(db: Session, season_id: int, rows: list[dict[str, str]]) -> ImportReport:
    """Plantilla: equipo;persona;rol (rol opcional). També entén capçaleres
    NOM I COGNOM / EQUIP / ENTRENADOR (l'entrenador es vincula com a coach)."""
    report = ImportReport()
    for i, row in enumerate(rows, start=2):
        team_name = _first(row, TEAM_HEADERS)
        person_name = _first(row, PERSON_HEADERS)
        role = _norm_role(_first(row, ROLE_HEADERS) or "player")
        coach_name = _norm(_first(row, COACH_HEADERS))
        # "entrenador" puede ser columna-flag (1/x/sí) en vez de un nombre
        if _truthy(coach_name) or coach_name.casefold() in ROLE_MAP:
            coach_name = ""
        if not team_name or not person_name:
            report.skipped += 1
            continue
        team = _get_or_create_team(
            db,
            season_id,
            team_name,
            row.get("categoria") or row.get("category"),
            report,
        )
        person = _get_or_create_person(db, season_id, person_name, role, report)
        if not team or not person:
            report.skipped += 1
            continue
        exists = (
            db.query(TeamMembership)
            .filter(
                TeamMembership.team_id == team.id,
                TeamMembership.person_id == person.id,
                TeamMembership.role == role,
            )
            .first()
        )
        if exists:
            report.skipped += 1
        else:
            db.add(TeamMembership(team_id=team.id, person_id=person.id, role=role))
            report.links_created += 1
        if coach_name:
            coach = _get_or_create_person(db, season_id, coach_name, "coach", report)
            if coach:
                exists_c = (
                    db.query(TeamMembership)
                    .filter(
                        TeamMembership.team_id == team.id,
                        TeamMembership.person_id == coach.id,
                        TeamMembership.role == "coach",
                    )
                    .first()
                )
                if exists_c:
                    report.skipped += 1
                else:
                    db.add(
                        TeamMembership(
                            team_id=team.id, person_id=coach.id, role="coach"
                        )
                    )
                    report.links_created += 1
    db.commit()
    return report


def import_teams_rows(db: Session, season_id: int, rows: list[dict[str, str]]) -> ImportReport:
    report = ImportReport()
    for i, row in enumerate(rows, start=2):
        name = _first(row, TEAM_HEADERS | {"nombre", "name", "nom"})
        category = row.get("categoria") or row.get("category") or ""
        if not _norm(name):
            report.skipped += 1
            continue
        before = report.teams_created
        _get_or_create_team(db, season_id, name, category, report)
        if report.teams_created == before:
            report.skipped += 1
    if report.teams_created:
        db.commit()
    else:
        db.rollback()
    return report


def import_people_rows(db: Session, season_id: int, rows: list[dict[str, str]]) -> ImportReport:
    report = ImportReport()
    for i, row in enumerate(rows, start=2):
        name = _first(row, PERSON_HEADERS)
        if not _norm(name):
            report.skipped += 1
            continue
        is_coach = _truthy(row.get("entrenador") or row.get("coach") or "")
        is_player = _truthy(row.get("jugador") or row.get("player") or "1")
        if is_coach and not (row.get("jugador") or row.get("player")):
            is_player = False
        role = "coach" if is_coach and not is_player else "player"
        team_name = _first(row, TEAM_HEADERS)
        before = report.people_created
        person = _get_or_create_person(db, season_id, name, role, report)
        if person:
            person.is_coach = person.is_coach or is_coach
            person.is_player = person.is_player or is_player or not is_coach
        linked_now = False
        if team_name and person:
            team = _get_or_create_team(
                db, season_id, team_name,
                row.get("categoria") or row.get("category"), report,
            )
            if team:
                exists = (
                    db.query(TeamMembership)
                    .filter(
                        TeamMembership.team_id == team.id,
                        TeamMembership.person_id == person.id,
                        TeamMembership.role == role,
                    )
                    .first()
                )
                if not exists:
                    db.add(
                        TeamMembership(
                            team_id=team.id, person_id=person.id, role=role
                        )
                    )
                    report.links_created += 1
                    linked_now = True
        if report.people_created == before and not linked_now:
            report.skipped += 1
    if report.people_created or report.links_created or report.teams_created:
        db.commit()
    else:
        db.rollback()
    return report


ROSTER_TEMPLATE = "equipo;persona;rol\nSenior A;Joan Garcia;jugador\nSenior A;Anna Coach;entrenador\nSenior B;Pere Lopez;jugador\n"
TEAMS_TEMPLATE = "equipo;categoria\nSenior A;Senior\nSenior B;Senior\n"
PEOPLE_TEMPLATE = "nombre;jugador;entrenador;equipo\nJoan Garcia;1;0;Senior A\nAnna Coach;0;1;Senior A\n"
